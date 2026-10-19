"""Testes das rotas do AlertSignal, no modo demonstração (banco temporário).

    python -m unittest discover -s tests -v
"""

import csv
import io
import os
import re
import sys
import tempfile
import unittest
from datetime import date, timedelta

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
os.environ['ALERTSIGNAL_DEMO'] = '1'

from werkzeug.security import check_password_hash   # noqa: E402

import database                      # noqa: E402
from demo_seed import popular, SENHA_DEMO   # noqa: E402
import app as aplicacao              # noqa: E402

PASTA = tempfile.mkdtemp(prefix='alertsignal-testes-')
database.usar_banco(os.path.join(PASTA, 'demo.db'))


def token(resp_html):
    m = re.search(r'name="csrf-token" content="([^"]+)"', resp_html)
    return m.group(1) if m else ''


class Base(unittest.TestCase):
    perfil = 'admin'

    def setUp(self):
        popular('completo')
        aplicacao.app.config['TESTING'] = True
        self.c = aplicacao.app.test_client()
        html = self.c.get('/login').get_data(as_text=True)
        self.csrf = token(html)
        r = self.c.post('/demo/entrar', data={'perfil': self.perfil, '_csrf': self.csrf})
        self.assertEqual(r.status_code, 302)
        # entrar limpa a sessão (e o token); as páginas seguintes trazem um novo
        self.csrf = token(self.c.get('/perfil').get_data(as_text=True))

    def post_json(self, url, dados=None, metodo='POST'):
        return self.c.open(url, method=metodo, json=dados or {}, headers={'X-CSRF-Token': self.csrf})


class TestPaginas(Base):
    def test_paginas_abrem_sem_erro(self):
        for url in ['/dashboard', '/empresas', '/empresa/1', '/empresa/6', '/empresa/16', '/relatorio',
                    '/relatorio?status=atencao', '/relatorio?status=vencidos&categoria=1', '/relatorio?empresa=3',
                    '/cadastros', '/cadastros?aba=inativos', '/responsaveis', '/historico', '/historico?page=2',
                    '/historico?page=abc', '/historico?page=-5', '/usuarios', '/configuracoes', '/perfil', '/demo/email']:
            with self.subTest(url=url):
                r = self.c.get(url)
                self.assertEqual(r.status_code, 200, url)

    def test_scripts_da_pagina_aparecem_uma_vez(self):
        html = self.c.get('/empresa/1').get_data(as_text=True)
        self.assertEqual(html.count('let linhaAtual'), 1)

    def test_erros_proprios(self):
        r = self.c.get('/nao-existe')
        self.assertEqual(r.status_code, 404)
        self.assertIn('Página não encontrada', r.get_data(as_text=True))
        self.assertEqual(self.c.get('/empresa/9999').status_code, 404)
        r = self.c.get('/api/doc/9999/editar')
        self.assertEqual(r.status_code, 405)
        self.assertEqual(r.get_json()['ok'], False)

    def test_apostrofo_nao_vai_para_javascript(self):
        html = self.c.get('/cadastros').get_data(as_text=True)
        self.assertNotIn("onsubmit=\"return confirm('", html)
        self.assertIn('Cantina D&#39;Ítalia', html)


class TestCsrf(Base):
    def test_post_sem_token_e_recusado(self):
        r = self.c.post('/api/doc/1/protocolo', json={'protocolo': 'X'})
        self.assertEqual(r.status_code, 400)
        r = self.c.post('/categoria/nova', data={'nome': 'Nova'})
        self.assertEqual(r.status_code, 302)
        conn = database.get_connection()
        self.assertIsNone(conn.execute("SELECT 1 FROM categorias WHERE nome='Nova'").fetchone())
        conn.close()

    def test_post_com_token_funciona(self):
        r = self.post_json('/api/doc/1/protocolo', {'protocolo': 'AVCB-TESTE'})
        self.assertTrue(r.get_json()['ok'])


class TestDocumentos(Base):
    def test_renovar_exige_data_futura(self):
        ontem = (date.today() - timedelta(days=1)).isoformat()
        self.assertEqual(self.post_json('/api/doc/1/renovar', {'vencimento': ontem}).status_code, 400)
        futuro = (date.today() + timedelta(days=200)).isoformat()
        self.assertTrue(self.post_json('/api/doc/1/renovar', {'vencimento': futuro}).get_json()['ok'])
        conn = database.get_connection()
        self.assertEqual(conn.execute('SELECT status FROM documentos WHERE id=1').fetchone()[0], 'OK')
        conn.close()

    def test_documento_sem_data_fica_sem_data(self):
        r = self.c.post('/cadastros/documento/novo', data={'empresa_id': 16, 'tipo': 'IPTU', '_csrf': self.csrf})
        self.assertEqual(r.status_code, 302)
        conn = database.get_connection()
        self.assertEqual(conn.execute("SELECT status FROM documentos WHERE empresa_id=16").fetchone()[0], 'NÃO TEM')
        conn.close()

    def test_historico_descreve_o_documento(self):
        self.post_json('/api/doc/1/protocolo', {'protocolo': 'P-1'})
        conn = database.get_connection()
        texto = conn.execute('SELECT descricao FROM historico ORDER BY id DESC LIMIT 1').fetchone()[0]
        conn.close()
        self.assertIn('Protocolo atualizado:', texto)
        self.assertNotIn('#', texto)

    def test_email_de_alerta(self):
        from notificacoes import montar_html, montar_texto
        itens = [{'empresa': "Cantina D'Ítalia <b>", 'documento': 'AVCB', 'vencimento': '2026-01-01', 'dias': -3, 'nivel': 'vencido'}]
        html = montar_html('Ana', itens, 'http://servidor:5000')
        self.assertIn('&lt;b&gt;', html)                      # nome escapado
        self.assertIn('Abrir no AlertSignal', html)
        self.assertIn('01/01/2026', html)
        self.assertNotIn('Abrir no AlertSignal', montar_html('Ana', itens))
        self.assertIn('Vencido há 3 dias', montar_texto('Ana', itens))

    def test_empresa_ordena_por_urgencia(self):
        html = self.c.get('/empresa/1').get_data(as_text=True)
        status = re.findall(r'data-status="([^"]+)"', html)
        ordem = {'VENCIDO': 0, 'RENOVAR': 1, 'NÃO TEM': 2, 'OK': 3}
        self.assertEqual(status, sorted(status, key=ordem.get))

    def test_alerta_nao_sai_no_demo(self):
        r = self.post_json('/api/testar-email').get_json()
        self.assertFalse(r['ok'])
        self.assertTrue(r.get('demo'))


class TestRegrasDaConferencia(Base):
    """Defeitos apontados na revisão de código final, para não voltarem."""

    def test_login_ignora_maiusculas(self):
        self.c.post('/usuarios/novo', data={'nome': 'Maria', 'email': 'Maria.Silva@Empresa.com', 'senha': 'abcdefgh',
                                            'nivel': 'visualizador', '_csrf': self.csrf})
        c2 = aplicacao.app.test_client()
        tok = token(c2.get('/login').get_data(as_text=True))
        r = c2.post('/login', data={'email': 'Maria.Silva@Empresa.com', 'senha': 'abcdefgh', '_csrf': tok})
        self.assertEqual(r.status_code, 302)

    def test_empresa_inativa_nao_gera_alerta(self):
        from notificacoes import buscar_alertas
        empresas = {i['empresa'] for a in buscar_alertas() for i in a['itens']}
        self.assertNotIn('Autopeças Avenida', empresas)

    def test_data_invalida_e_recusada(self):
        self.assertTrue(self.post_json('/api/doc/1/editar', {'vencimento': '2030-05-20'}).get_json()['ok'])
        r = self.post_json('/api/doc/1/editar', {'vencimento': '31/12/2027'})
        self.assertEqual(r.status_code, 400)
        conn = database.get_connection()
        self.assertEqual(conn.execute('SELECT vencimento FROM documentos WHERE id=1').fetchone()[0], '2030-05-20')
        conn.close()

    def test_remover_responsavel_apaga_vinculos(self):
        self.c.post('/responsaveis/1/excluir', data={'_csrf': self.csrf})
        conn = database.get_connection()
        self.assertEqual(conn.execute('SELECT COUNT(*) FROM documento_responsavel WHERE responsavel_id=1').fetchone()[0], 0)
        conn.close()

    def test_ficha_de_empresa_inativa_exporta(self):
        r = self.c.get('/exportar?fmt=csv&empresa=17')
        self.assertGreater(len(r.data.decode('utf-8-sig').strip().splitlines()), 1)


class TestExportacao(Base):
    def test_excel(self):
        from openpyxl import load_workbook
        r = self.c.get('/exportar?fmt=xlsx&status=atencao')
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.data)).active
        self.assertEqual(ws['A5'].value, 'Empresa')
        self.assertIn('Vencidos:', ws['A3'].value)
        self.assertNotEqual(ws['A7'].fill.fgColor.rgb, '00111111')   # zebra não pode ser preta
        self.assertEqual(ws['F6'].number_format, 'DD/MM/YYYY')
        self.assertEqual(ws['G5'].value, 'Dias para vencer')

    def test_csv_com_ponto_e_virgula(self):
        r = self.c.get('/exportar?fmt=csv&empresa=6')
        texto = r.data.decode('utf-8-sig')
        linhas = list(csv.reader(io.StringIO(texto), delimiter=';'))
        self.assertEqual(linhas[0][0], 'Empresa')
        self.assertTrue(all(len(l) == len(linhas[0]) for l in linhas))
        self.assertTrue(all(l[0] == "Cantina D'Ítalia" for l in linhas[1:]))


class TestSeguranca(Base):
    """Senha inicial, limite de tentativas, cabeçalhos e exportação sem fórmula."""

    def setUp(self):
        aplicacao._falhas.clear()
        super().setUp()

    def tearDown(self):
        aplicacao._falhas.clear()
        for chave in ('ADMIN_EMAIL', 'ADMIN_SENHA'):
            os.environ.pop(chave, None)

    def entrar(self, cliente, email, senha):
        tok = token(cliente.get('/login').get_data(as_text=True))
        return cliente.post('/login', data={'email': email, 'senha': senha, '_csrf': tok})

    def test_cabecalhos_de_seguranca(self):
        r = self.c.get('/dashboard')
        self.assertIn("default-src 'self'", r.headers['Content-Security-Policy'])
        self.assertIn("frame-ancestors 'none'", r.headers['Content-Security-Policy'])
        self.assertEqual(r.headers['X-Content-Type-Options'], 'nosniff')
        self.assertEqual(r.headers['X-Frame-Options'], 'DENY')
        self.assertEqual(r.headers['Referrer-Policy'], 'same-origin')

    def test_pagina_nao_carrega_nada_de_fora(self):
        html = self.c.get('/dashboard').get_data(as_text=True)
        self.assertNotIn('fonts.googleapis.com', html)
        self.assertNotIn('cdn.jsdelivr.net', html)

    def test_login_bloqueia_depois_de_cinco_erros(self):
        c2 = aplicacao.app.test_client()
        for i in range(aplicacao.MAX_FALHAS_EMAIL):
            self.assertEqual(self.entrar(c2, 'admin@alertsignal.com', f'errada{i}').status_code, 200)
        r = self.entrar(c2, 'admin@alertsignal.com', SENHA_DEMO)     # nem a senha certa entra
        self.assertEqual(r.status_code, 429)
        self.assertIn('Muitas tentativas', r.get_data(as_text=True))
        # o bloqueio é do e-mail: outra conta do mesmo computador entra
        self.assertEqual(self.entrar(c2, 'visualizador@alertsignal.com', SENHA_DEMO).status_code, 302)

    def test_acerto_zera_a_contagem(self):
        c2 = aplicacao.app.test_client()
        for rodada in range(2):
            for i in range(aplicacao.MAX_FALHAS_EMAIL - 1):
                self.entrar(c2, 'admin@alertsignal.com', 'errada')
            self.assertEqual(self.entrar(c2, 'admin@alertsignal.com', SENHA_DEMO).status_code, 302, rodada)

    def test_admin_inicial_tem_senha_aleatoria(self):
        original = database.DB_PATH
        vazio = os.path.join(PASTA, 'vazio.db')
        if os.path.exists(vazio):
            os.remove(vazio)
        database.usar_banco(vazio)
        try:
            database.init_db()
            email, gerada = aplicacao.criar_admin_padrao()
            self.assertEqual(email, 'admin@example.com')
            self.assertGreaterEqual(len(gerada), 12)
            conn = database.get_connection()
            senha = conn.execute('SELECT senha FROM usuarios').fetchone()[0]
            conn.close()
            self.assertTrue(check_password_hash(senha, gerada))
            self.assertIsNone(aplicacao.criar_admin_padrao())     # com usuário, não cria outro
        finally:
            database.usar_banco(original)

    def test_admin_inicial_vem_do_env(self):
        original = database.DB_PATH
        vazio = os.path.join(PASTA, 'vazio-env.db')
        if os.path.exists(vazio):
            os.remove(vazio)
        database.usar_banco(vazio)
        os.environ['ADMIN_EMAIL'] = 'Chefe@Empresa.com'
        os.environ['ADMIN_SENHA'] = 'uma-senha-longa'
        try:
            database.init_db()
            self.assertEqual(aplicacao.criar_admin_padrao(), ('chefe@empresa.com', None))
        finally:
            database.usar_banco(original)

    def test_exportacao_nao_vira_formula(self):
        from openpyxl import load_workbook
        conn = database.get_connection()
        conn.execute('UPDATE documentos SET observacoes=?, protocolo=? WHERE id=1',
                     ('=HYPERLINK("http://exemplo.invalid","clique")', '+55 71'))
        conn.commit()
        conn.close()
        ws = load_workbook(io.BytesIO(self.c.get('/exportar?fmt=xlsx').data)).active
        tipos = {c.data_type for linha in ws.iter_rows(min_row=6) for c in linha}
        self.assertNotIn('f', tipos)
        textos = [c.value for linha in ws.iter_rows(min_row=6) for c in linha if isinstance(c.value, str)]
        self.assertIn('=HYPERLINK("http://exemplo.invalid","clique")', textos)

        texto = self.c.get('/exportar?fmt=csv').data.decode('utf-8-sig')
        linhas = list(csv.reader(io.StringIO(texto), delimiter=';'))
        celulas = [c for l in linhas for c in l]
        self.assertIn('\'=HYPERLINK("http://exemplo.invalid","clique")', celulas)
        self.assertIn("'+55 71", celulas)
        # os dias negativos (documentos vencidos) continuam números
        self.assertTrue(any(re.fullmatch(r'-\d+', l[6]) for l in linhas[1:]))


class TestVisualizador(Base):
    perfil = 'visualizador'

    def test_consulta_funciona(self):
        for url in ['/dashboard', '/empresas', '/empresa/1', '/relatorio', '/cadastros', '/responsaveis', '/historico', '/perfil']:
            with self.subTest(url=url):
                self.assertEqual(self.c.get(url).status_code, 200)

    def test_nao_ve_acoes_de_admin(self):
        html = self.c.get('/empresa/1').get_data(as_text=True)
        self.assertNotIn('data-acao="editar"', html)
        self.assertNotIn('id="modal-editar"', html)

    def test_servidor_bloqueia_alteracoes(self):
        for url, metodo in [('/api/doc/1/protocolo', 'POST'), ('/api/doc/1/renovar', 'POST'),
                            ('/api/doc/1/responsavel', 'POST'), ('/api/doc/1/responsavel/1', 'DELETE'),
                            ('/api/doc/1/editar', 'POST'), ('/api/doc/1/excluir', 'DELETE'),
                            ('/api/testar-email', 'POST')]:
            with self.subTest(url=url):
                self.assertEqual(self.post_json(url, {'vencimento': '2099-01-01'}, metodo).status_code, 403)
        self.assertEqual(self.c.get('/usuarios').status_code, 302)



def integridade_ok(teste):
    """Nenhuma referência quebrada no banco (PRAGMA foreign_key_check vazio)."""
    conn = database.get_connection()
    problemas = conn.execute('PRAGMA foreign_key_check').fetchall()
    conn.close()
    teste.assertEqual([tuple(p) for p in problemas], [])


class TestIntegridade(Base):
    """Chaves estrangeiras ligadas, senha mínima e CSRF em toda rota de escrita."""

    def test_chaves_estrangeiras_ligadas(self):
        conn = database.get_connection()
        self.assertEqual(conn.execute('PRAGMA foreign_keys').fetchone()[0], 1)
        conn.close()
        integridade_ok(self)

    def test_excluir_empresa_com_historico(self):
        self.c.post('/empresa/nova', data={'nome': 'Empresa Teste', 'categoria_id': 1, '_csrf': self.csrf})
        conn = database.get_connection()
        emp_id = conn.execute("SELECT id FROM empresas WHERE nome='Empresa Teste'").fetchone()[0]
        conn.close()
        r = self.c.post(f'/empresa/{emp_id}/excluir', data={'_csrf': self.csrf})
        self.assertEqual(r.status_code, 302)
        conn = database.get_connection()
        self.assertIsNone(conn.execute('SELECT 1 FROM empresas WHERE id=?', (emp_id,)).fetchone())
        self.assertIsNotNone(conn.execute(
            "SELECT 1 FROM historico WHERE descricao='Empresa cadastrada: Empresa Teste'").fetchone())
        conn.close()
        integridade_ok(self)

    def test_excluir_documento_com_historico(self):
        self.post_json('/api/doc/1/protocolo', {'protocolo': 'P-1'})
        self.assertTrue(self.post_json('/api/doc/1/excluir', metodo='DELETE').get_json()['ok'])
        integridade_ok(self)

    def test_excluir_usuario_com_historico(self):
        conn = database.get_connection()
        conn.execute("INSERT INTO historico (tipo, descricao, usuario_id) VALUES ('tramite', 'feito pela Marina', 2)")
        conn.commit()
        conn.close()
        self.assertEqual(self.c.post('/usuarios/2/excluir', data={'_csrf': self.csrf}).status_code, 302)
        conn = database.get_connection()
        self.assertIsNone(conn.execute('SELECT 1 FROM usuarios WHERE id=2').fetchone())
        conn.close()
        integridade_ok(self)

    def test_senha_curta_e_recusada(self):
        self.c.post('/usuarios/novo', data={'nome': 'Curta', 'email': 'curta@example.com', 'senha': '1234567',
                                            'nivel': 'visualizador', '_csrf': self.csrf})
        self.c.post('/usuarios/2/redefinir-senha', data={'senha_nova': '1234567', '_csrf': self.csrf})
        conn = database.get_connection()
        self.assertIsNone(conn.execute("SELECT 1 FROM usuarios WHERE email='curta@example.com'").fetchone())
        senha_b = conn.execute('SELECT senha FROM usuarios WHERE id=2').fetchone()[0]
        conn.close()
        self.assertFalse(check_password_hash(senha_b, '1234567'))
        self.assertIn('minlength="8"', self.c.get('/perfil').get_data(as_text=True))

    def test_reiniciar_demo_exige_token(self):
        self.post_json('/api/doc/1/protocolo', {'protocolo': 'MARCA'})
        self.c.post('/demo/reiniciar', data={'cenario': 'vazio'})       # sem _csrf e sem Origin
        conn = database.get_connection()
        self.assertEqual(conn.execute('SELECT protocolo FROM documentos WHERE id=1').fetchone()[0], 'MARCA')
        conn.close()

    def test_faxina_das_falhas_antigas(self):
        aplicacao._falhas.clear()
        antigo = aplicacao.time.monotonic() - aplicacao.JANELA_FALHAS - 1
        aplicacao._falhas['e:sumiu@example.com'] = [antigo]
        aplicacao._ultima_faxina[0] = antigo
        try:
            with aplicacao.app.test_request_context('/login', method='POST'):
                aplicacao._registrar_falha('outro@example.com')
            self.assertNotIn('e:sumiu@example.com', aplicacao._falhas)
            self.assertIn('e:outro@example.com', aplicacao._falhas)
        finally:
            aplicacao._falhas.clear()


class TestSessaoRevalidada(Base):
    """A sessão vale o que o banco diz agora, não o que dizia no login."""

    def entrar_como(self, email, senha):
        cliente = aplicacao.app.test_client()
        tok = token(cliente.get('/login').get_data(as_text=True))
        r = cliente.post('/login', data={'email': email, 'senha': senha, '_csrf': tok})
        self.assertEqual(r.status_code, 302)
        return cliente, token(cliente.get('/perfil').get_data(as_text=True))

    def criar_admin_b(self):
        self.c.post('/usuarios/novo', data={'nome': 'Admin B', 'email': 'admin.b@example.com', 'senha': 'senha-do-b',
                                            'nivel': 'admin', '_csrf': self.csrf})
        conn = database.get_connection()
        uid = conn.execute("SELECT id FROM usuarios WHERE email='admin.b@example.com'").fetchone()[0]
        conn.close()
        return uid

    def test_usuario_excluido_perde_o_acesso(self):
        # O caso da avaliação: o admin A exclui o admin B, que estava conectado
        uid = self.criar_admin_b()
        b, tok_b = self.entrar_como('admin.b@example.com', 'senha-do-b')
        self.assertEqual(b.get('/usuarios').status_code, 200)
        self.assertEqual(self.c.post(f'/usuarios/{uid}/excluir', data={'_csrf': self.csrf}).status_code, 302)
        r = b.get('/usuarios')
        self.assertEqual(r.status_code, 302)
        self.assertIn('/login', r.headers['Location'])
        self.assertEqual(b.get('/configuracoes').status_code, 302)
        b.post('/categoria/nova', data={'nome': 'Criada por excluido', '_csrf': tok_b})
        r = b.post('/api/doc/1/protocolo', json={'protocolo': 'X'}, headers={'X-CSRF-Token': tok_b})
        self.assertIn(r.status_code, (400, 401))
        conn = database.get_connection()
        self.assertIsNone(conn.execute("SELECT 1 FROM categorias WHERE nome='Criada por excluido'").fetchone())
        self.assertNotEqual(conn.execute('SELECT protocolo FROM documentos WHERE id=1').fetchone()[0], 'X')
        conn.close()

    def test_api_responde_401_para_sessao_revogada(self):
        uid = self.criar_admin_b()
        b, tok_b = self.entrar_como('admin.b@example.com', 'senha-do-b')
        conn = database.get_connection()
        conn.execute('UPDATE usuarios SET ativo=0 WHERE id=?', (uid,))
        conn.commit()
        conn.close()
        # o token CSRF continua válido; quem barra é a revalidação do usuário
        r = b.post('/api/doc/1/protocolo', json={'protocolo': 'X'}, headers={'X-CSRF-Token': tok_b})
        self.assertEqual(r.status_code, 401)

    def test_usuario_desativado_perde_o_acesso(self):
        uid = self.criar_admin_b()
        b, _ = self.entrar_como('admin.b@example.com', 'senha-do-b')
        conn = database.get_connection()
        conn.execute('UPDATE usuarios SET ativo=0 WHERE id=?', (uid,))
        conn.commit()
        conn.close()
        self.assertEqual(b.get('/dashboard').status_code, 302)

    def test_nivel_vem_do_banco(self):
        uid = self.criar_admin_b()
        b, tok_b = self.entrar_como('admin.b@example.com', 'senha-do-b')
        conn = database.get_connection()
        conn.execute("UPDATE usuarios SET nivel='visualizador' WHERE id=?", (uid,))
        conn.commit()
        conn.close()
        self.assertEqual(b.get('/usuarios').status_code, 302)          # volta para a visão geral
        self.assertEqual(b.get('/dashboard').status_code, 200)         # continua podendo consultar
        r = b.post('/api/doc/1/protocolo', json={'protocolo': 'X'}, headers={'X-CSRF-Token': tok_b})
        self.assertEqual(r.status_code, 403)

    def test_redefinir_senha_derruba_sessoes_antigas(self):
        uid = self.criar_admin_b()
        b, _ = self.entrar_como('admin.b@example.com', 'senha-do-b')
        self.c.post(f'/usuarios/{uid}/redefinir-senha', data={'senha_nova': 'senha-nova-b', '_csrf': self.csrf})
        self.assertEqual(b.get('/dashboard').status_code, 302)
        self.assertEqual(self.c.get('/usuarios').status_code, 200)     # quem redefiniu continua
        self.entrar_como('admin.b@example.com', 'senha-nova-b')

    def test_trocar_a_propria_senha_mantem_esta_sessao(self):
        self.criar_admin_b()
        b1, tok1 = self.entrar_como('admin.b@example.com', 'senha-do-b')
        b2, _ = self.entrar_como('admin.b@example.com', 'senha-do-b')
        r = b1.post('/perfil', data={'acao': 'senha', 'senha_atual': 'senha-do-b', 'senha_nova': 'outra-senha-b',
                                     'senha_conf': 'outra-senha-b', '_csrf': tok1})
        self.assertIn('Senha alterada', r.get_data(as_text=True))
        self.assertEqual(b1.get('/dashboard').status_code, 200)
        self.assertEqual(b2.get('/dashboard').status_code, 302)


if __name__ == '__main__':
    unittest.main()
