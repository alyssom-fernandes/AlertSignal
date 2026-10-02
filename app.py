from flask import Flask, render_template, request, redirect, url_for, session, jsonify, flash, abort
from werkzeug.exceptions import HTTPException
from werkzeug.security import generate_password_hash, check_password_hash
from apscheduler.schedulers.background import BackgroundScheduler
from datetime import date, datetime, timedelta
from functools import wraps
from dotenv import load_dotenv
import hmac
import math
import os
import re
import secrets
import sys
import atexit

load_dotenv()

# Modo demonstração: banco separado (demo.db) com dados fictícios, entrada
# sem senha e nenhum e-mail enviado. Liga com `python app.py --demo`.
MODO_DEMO = '--demo' in sys.argv or os.environ.get('ALERTSIGNAL_DEMO') == '1'
if MODO_DEMO:
    os.environ['ALERTSIGNAL_DEMO'] = '1'

import database
from database import get_connection, init_db, inserir_configuracoes_padrao, get_config, set_config
from importar_planilha import importar
from notificacoes import executar_verificacao_diaria, calcular_dias, recalcular_status

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY') or os.urandom(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax')
scheduler = None

TIPOS_DOC = [
    'AVCB', 'Licença de Operação', 'LAC - Licença de Transportes',
    'Alvará Municipal', 'Alvará Sanitário', 'FEASPOL',
    'CADASTUR', 'IPTU', 'Alvará Policial',
]

NIVEIS = {'admin': 'Administrador', 'visualizador': 'Visualizador'}

# Filtros de status do relatório e da exportação
FILTROS_STATUS = {
    'todos':    ('Todos os documentos', None),
    'atencao':  ('Vencidos e a renovar', ('VENCIDO', 'RENOVAR')),
    'vencidos': ('Só vencidos',          ('VENCIDO',)),
    'renovar':  ('Só a renovar',         ('RENOVAR',)),
    'em_dia':   ('Só em dia',            ('OK',)),
    'sem_data': ('Sem data cadastrada',  ('NÃO TEM',)),
}

MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
MESES_EXTENSO = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho',
                 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']
DIAS_SEMANA = ['segunda-feira', 'terça-feira', 'quarta-feira', 'quinta-feira',
               'sexta-feira', 'sábado', 'domingo']

# ─── PROTEÇÃO CONTRA CSRF ─────────────────────────────────────────────────────
# Todo formulário leva um campo _csrf e todo fetch que altera dados leva o
# cabeçalho X-CSRF-Token (ver static/js/app.js). Sem dependência externa.

def csrf_token():
    if '_csrf' not in session:
        session['_csrf'] = secrets.token_urlsafe(32)
    return session['_csrf']

app.jinja_env.globals['csrf_token'] = csrf_token

@app.before_request
def verificar_csrf():
    if request.method not in ('POST', 'PUT', 'PATCH', 'DELETE'):
        return
    if MODO_DEMO and request.path == '/demo/reiniciar' and not request.headers.get('Origin'):
        return   # ferramenta de fotos (sem navegador); formulários do navegador mandam o token
    enviado  = request.form.get('_csrf') or request.headers.get('X-CSRF-Token') or ''
    esperado = session.get('_csrf') or ''
    if esperado and hmac.compare_digest(enviado, esperado):
        return
    if request.path.startswith('/api/'):
        return jsonify({'ok': False, 'erro': 'A sessão expirou. Recarregue a página e tente de novo.'}), 400
    flash('A sessão expirou. Tente de novo.', 'erro')
    return redirect(request.referrer or url_for('login'))

# ─── DECORADORES ──────────────────────────────────────────────────────────────

def _eh_api():
    return request.path.startswith('/api/')

def login_obrigatorio(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if 'usuario_id' not in session:
            if _eh_api():
                return jsonify({'ok': False, 'erro': 'Entre de novo para continuar.'}), 401
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return wrapper

def admin_obrigatorio(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if 'usuario_id' not in session:
            if _eh_api():
                return jsonify({'ok': False, 'erro': 'Entre de novo para continuar.'}), 401
            return redirect(url_for('login'))
        if session.get('usuario_nivel') != 'admin':
            if _eh_api():
                return jsonify({'ok': False, 'erro': 'Ação restrita a administradores.'}), 403
            flash('Acesso restrito a administradores.', 'erro')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return wrapper

# ─── HELPERS ──────────────────────────────────────────────────────────────────

def criar_admin_padrao():
    conn = get_connection()
    total = conn.execute('SELECT COUNT(*) FROM usuarios').fetchone()[0]
    if total == 0:
        conn.execute('INSERT INTO usuarios (nome, email, senha, nivel) VALUES (?,?,?,?)',
                     ('Administrador', 'admin@grupozen.com.br', generate_password_hash('zen2024'), 'admin'))
        conn.commit()
    conn.close()

def status_badge(status):
    return {'OK':'ok','RENOVAR':'warn','VENCIDO':'danger','NÃO TEM':'none'}.get(status,'none')

def registrar_historico(conn, descricao, tipo='tramite', empresa_id=None, documento_id=None):
    uid = session.get('usuario_id')
    conn.execute('INSERT INTO historico (tipo, descricao, empresa_id, documento_id, usuario_id) VALUES (?,?,?,?,?)',
                 (tipo, descricao, empresa_id, documento_id, uid))

def dados_documento(conn, doc_id):
    """Tipo do documento e nome/id da empresa, para mensagens e histórico."""
    return conn.execute('''SELECT d.id, d.tipo, e.nome AS empresa, e.id AS empresa_id
                           FROM documentos d JOIN empresas e ON e.id=d.empresa_id
                           WHERE d.id=?''', (doc_id,)).fetchone()

def recalcular_todos(conn):
    docs = conn.execute('SELECT id, vencimento, status FROM documentos').fetchall()
    limite = int(get_config('alerta_dias_30') or 30)
    for doc in docs:
        novo = recalcular_status(calcular_dias(doc['vencimento']), limite)
        if novo != doc['status']:
            conn.execute('UPDATE documentos SET status=? WHERE id=?', (novo, doc['id']))
    conn.commit()

def get_stats(conn):
    rows = conn.execute('''SELECT d.status, COUNT(*) AS total FROM documentos d
                           JOIN empresas e ON e.id=d.empresa_id
                           WHERE e.ativa=1 GROUP BY d.status''').fetchall()
    stats = {'OK':0,'RENOVAR':0,'VENCIDO':0,'NÃO TEM':0}
    for r in rows:
        stats[r['status']] = r['total']
    return stats

def data_valida(valor):
    """Aceita só datas ISO (AAAA-MM-DD) de verdade; devolve None para o resto."""
    if not valor:
        return None
    try:
        return datetime.strptime(valor, '%Y-%m-%d').strftime('%Y-%m-%d')
    except (TypeError, ValueError):
        return None

def ler_filtros():
    status = request.args.get('status', 'todos')
    if status not in FILTROS_STATUS:
        status = 'todos'
    return {
        'status': status,
        'categoria': request.args.get('categoria', type=int),
        'empresa': request.args.get('empresa', type=int),
    }

def consultar_documentos(conn, filtros):
    """Documentos de empresas ativas, com responsáveis, segundo os filtros."""
    where, params = ([] if filtros.get('empresa') else ['e.ativa=1']), []
    situacoes = FILTROS_STATUS[filtros['status']][1]
    if situacoes:
        where.append(f"d.status IN ({','.join('?' * len(situacoes))})")
        params += list(situacoes)
    if filtros.get('categoria'):
        where.append('e.categoria_id=?'); params.append(filtros['categoria'])
    if filtros.get('empresa'):
        where.append('e.id=?'); params.append(filtros['empresa'])
    linhas = conn.execute(f'''
        SELECT d.id, d.tipo, d.protocolo, d.vencimento, d.status, d.observacoes,
               e.id AS empresa_id, e.nome AS empresa, e.cnpj, c.nome AS categoria,
               (SELECT GROUP_CONCAT(r.nome, ', ') FROM documento_responsavel dr
                  JOIN responsaveis r ON r.id=dr.responsavel_id
                 WHERE dr.documento_id=d.id AND r.ativo=1) AS responsaveis
        FROM documentos d
        JOIN empresas e ON e.id=d.empresa_id
        JOIN categorias c ON c.id=e.categoria_id
        WHERE {' AND '.join(where) or '1=1'}
        ORDER BY c.nome, e.nome, d.vencimento IS NULL, d.vencimento, d.tipo
    ''', params).fetchall()
    docs = []
    for l in linhas:
        d = dict(l)
        d['dias'] = calcular_dias(d['vencimento'])
        docs.append(d)
    return docs

def descrever_filtros(conn, filtros):
    partes = [FILTROS_STATUS[filtros['status']][0]]
    if filtros.get('categoria'):
        r = conn.execute('SELECT nome FROM categorias WHERE id=?', (filtros['categoria'],)).fetchone()
        if r: partes.append(f"categoria {r['nome']}")
    if filtros.get('empresa'):
        r = conn.execute('SELECT nome FROM empresas WHERE id=?', (filtros['empresa'],)).fetchone()
        if r: partes.append(r['nome'])
    return ' · '.join(partes)

def _slug(texto):
    import unicodedata
    t = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', t.lower()).strip('-')[:40] or 'documentos'

# ─── CONTEXT PROCESSOR ────────────────────────────────────────────────────────

@app.context_processor
def inject_globals():
    urgentes_count = 0
    if 'usuario_id' in session:
        try:
            conn = get_connection()
            urgentes_count = conn.execute(
                "SELECT COUNT(*) FROM documentos d JOIN empresas e ON e.id=d.empresa_id "
                "WHERE d.status IN ('VENCIDO','RENOVAR') AND e.ativa=1"
            ).fetchone()[0]
            conn.close()
        except Exception:
            pass
    try:
        limite_renovar = int(get_config('alerta_dias_30') or 30)
    except (TypeError, ValueError):
        limite_renovar = 30
    return {'now': datetime.now(), 'urgentes_count': urgentes_count, 'modo_demo': MODO_DEMO,
            'limite_renovar': limite_renovar,
            'eh_admin': session.get('usuario_nivel') == 'admin', 'niveis': NIVEIS}

# ─── FILTROS JINJA ────────────────────────────────────────────────────────────

@app.template_filter('formatar_data')
def formatar_data(valor):
    if not valor: return '—'
    try: return datetime.strptime(valor, '%Y-%m-%d').strftime('%d/%m/%Y')
    except (TypeError, ValueError): return valor

@app.template_filter('formatar_data_hora')
def formatar_data_hora(valor):
    if not valor: return '—'
    try: return datetime.strptime(valor[:16], '%Y-%m-%d %H:%M').strftime('%d/%m/%Y às %H:%M')
    except (TypeError, ValueError): return valor

@app.template_filter('data_extenso')
def data_extenso(d):
    return f'{DIAS_SEMANA[d.weekday()]}, {d.day} de {MESES_EXTENSO[d.month - 1]} de {d.year}'

@app.template_filter('dia_rotulo')
def dia_rotulo(valor):
    """'2026-10-01 17:31:00' → 'Hoje', 'Ontem' ou '29 de setembro de 2026'."""
    try:
        d = datetime.strptime(valor[:10], '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return valor
    hoje = date.today()
    if d == hoje: return 'Hoje'
    if d == hoje - timedelta(days=1): return 'Ontem'
    return f'{d.day} de {MESES_EXTENSO[d.month - 1]} de {d.year}'

@app.template_filter('status_classe')
def status_classe(status):
    return status_badge(status)

@app.template_filter('status_rotulo')
def status_rotulo(status):
    return {'OK': 'Em dia', 'RENOVAR': 'Renovar', 'VENCIDO': 'Vencido', 'NÃO TEM': 'Sem data'}.get(status, status)

@app.template_filter('dias_texto')
def dias_texto(dias):
    if dias is None: return '—'
    if dias == -1: return 'ontem'
    if dias < 0: return f'há {abs(dias)} dias'
    if dias == 0: return 'hoje'
    if dias == 1: return 'amanhã'
    return f'em {dias} dias'

@app.template_filter('prazo_texto')
def prazo_texto(dias):
    if dias is None: return 'Sem data'
    if dias < -1: return f'Vencido há {abs(dias)} dias'
    if dias == -1: return 'Venceu ontem'
    if dias == 0: return 'Vence hoje'
    if dias == 1: return 'Vence amanhã'
    return f'Vence em {dias} dias'

# ─── ERROS ────────────────────────────────────────────────────────────────────

MENSAGENS_ERRO = {
    403: ('Acesso negado', 'Você não tem permissão para abrir esta página.'),
    404: ('Página não encontrada', 'O endereço pode ter mudado ou o registro foi removido.'),
    405: ('Ação não permitida', 'Esta página não aceita esse tipo de acesso.'),
    500: ('Algo deu errado', 'O sistema encontrou um erro inesperado. Tente de novo em instantes.'),
}

def _pagina_erro(codigo):
    titulo, texto = MENSAGENS_ERRO.get(codigo, ('Não foi possível concluir', 'O pedido não pôde ser atendido. Volte e tente de novo.'))
    if _eh_api():
        return jsonify({'ok': False, 'erro': titulo}), codigo
    return render_template('erro.html', codigo=codigo, titulo=titulo, texto=texto), codigo

@app.errorhandler(HTTPException)
def erro_http(e):
    return _pagina_erro(e.code or 500)

@app.errorhandler(Exception)
def erro_inesperado(e):
    app.logger.exception('Erro inesperado em %s', request.path)
    return _pagina_erro(500)

# ─── AUTH ──────────────────────────────────────────────────────────────────────

@app.route('/favicon.ico')
def favicon():
    return app.send_static_file('favicon.ico')

@app.route('/')
def index():
    if 'usuario_id' in session: return redirect(url_for('dashboard'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET','POST'])
def login():
    erro = None
    email = ''
    if request.method == 'POST':
        email = request.form.get('email','').strip()
        senha = request.form.get('senha','')
        conn = get_connection()
        u = conn.execute('SELECT * FROM usuarios WHERE lower(email)=? AND ativo=1', (email.lower(),)).fetchone()
        conn.close()
        if u and check_password_hash(u['senha'], senha):
            session.clear()
            session['usuario_id']    = u['id']
            session['usuario_nome']  = u['nome']
            session['usuario_nivel'] = u['nivel']
            return redirect(url_for('dashboard'))
        erro = 'E-mail ou senha incorretos.'
    return render_template('login.html', erro=erro, email=email)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

# ─── DEMONSTRAÇÃO ─────────────────────────────────────────────────────────────

@app.route('/demo/entrar', methods=['POST'])
def demo_entrar():
    if not MODO_DEMO: abort(404)
    from demo_seed import USUARIOS
    perfil = request.form.get('perfil', 'admin')
    email = next((e for _, e, nivel in USUARIOS if nivel == perfil), USUARIOS[0][1])
    conn = get_connection()
    u = conn.execute('SELECT * FROM usuarios WHERE email=?', (email,)).fetchone()
    conn.close()
    if not u: abort(404)
    session.clear()
    session['usuario_id']    = u['id']
    session['usuario_nome']  = u['nome']
    session['usuario_nivel'] = u['nivel']
    return redirect(url_for('dashboard'))

@app.route('/demo/email')
def demo_email():
    """Mostra o e-mail de alerta que o primeiro responsável receberia hoje."""
    if not MODO_DEMO: abort(404)
    from notificacoes import buscar_alertas, montar_html
    alertas = buscar_alertas()
    url = request.host_url.rstrip('/')
    if not alertas: return montar_html('Responsável', [], url)
    return montar_html(alertas[0]['nome'], alertas[0]['itens'], url)

@app.route('/demo/erro')
def demo_erro():
    """Força um erro interno para conferir a página 500 (só na demonstração)."""
    if not MODO_DEMO: abort(404)
    raise RuntimeError('Erro provocado de propósito para conferir a página 500.')

@app.route('/demo/reiniciar', methods=['POST'])
def demo_reiniciar():
    if not MODO_DEMO: abort(404)
    from demo_seed import popular
    cenario = request.form.get('cenario', 'completo')
    popular('vazio' if cenario == 'vazio' else 'completo')
    if 'usuario_id' in session:
        flash('Dados de demonstração recriados.', 'ok')
    return redirect(request.referrer or url_for('login'))

# ─── PERFIL ────────────────────────────────────────────────────────────────────

@app.route('/perfil', methods=['GET','POST'])
@login_obrigatorio
def perfil():
    conn = get_connection()
    u = conn.execute('SELECT * FROM usuarios WHERE id=?', (session['usuario_id'],)).fetchone()
    if not u:
        conn.close(); session.clear(); return redirect(url_for('login'))
    erro = ok = None
    if request.method == 'POST':
        acao = request.form.get('acao')
        if acao == 'nome':
            novo_nome = request.form.get('nome','').strip()
            if novo_nome:
                conn.execute('UPDATE usuarios SET nome=? WHERE id=?', (novo_nome, session['usuario_id']))
                conn.commit()
                session['usuario_nome'] = novo_nome
                ok = 'Nome atualizado.'
            else:
                erro = 'Informe um nome.'
        elif acao == 'senha':
            atual = request.form.get('senha_atual','')
            nova  = request.form.get('senha_nova','')
            conf  = request.form.get('senha_conf','')
            if not check_password_hash(u['senha'], atual):
                erro = 'Senha atual incorreta.'
            elif len(nova) < 6:
                erro = 'A nova senha precisa ter pelo menos 6 caracteres.'
            elif nova != conf:
                erro = 'As senhas não coincidem.'
            else:
                conn.execute('UPDATE usuarios SET senha=? WHERE id=?', (generate_password_hash(nova), session['usuario_id']))
                conn.commit()
                ok = 'Senha alterada com sucesso.'
        u = conn.execute('SELECT * FROM usuarios WHERE id=?', (session['usuario_id'],)).fetchone()
    conn.close()
    return render_template('perfil.html', usuario=u, erro=erro, ok=ok)

# ─── DASHBOARD ────────────────────────────────────────────────────────────────

def vencimentos_por_mes(conn, hoje):
    """Documentos (empresas ativas) que vencem em cada um dos próximos 12 meses."""
    inicio = hoje.replace(day=1)
    meses = []
    for i in range(12):
        ano, mes = inicio.year + (inicio.month - 1 + i) // 12, (inicio.month - 1 + i) % 12 + 1
        meses.append({'chave': f'{ano:04d}-{mes:02d}', 'rotulo': MESES[mes - 1], 'ano': ano,
                      'extenso': f'{MESES_EXTENSO[mes - 1]} de {ano}', 'total': 0, 'urgentes': 0})
    por_chave = {m['chave']: m for m in meses}
    linhas = conn.execute('''SELECT d.vencimento FROM documentos d JOIN empresas e ON e.id=d.empresa_id
                             WHERE e.ativa=1 AND d.vencimento >= ?''', (hoje.isoformat(),)).fetchall()
    limite_renovar = int(get_config('alerta_dias_30') or 30)
    for l in linhas:
        m = por_chave.get(l['vencimento'][:7])
        if m:
            m['total'] += 1
            if calcular_dias(l['vencimento']) <= limite_renovar:
                m['urgentes'] += 1
    return meses

@app.route('/dashboard')
@login_obrigatorio
def dashboard():
    conn = get_connection()
    recalcular_todos(conn)
    stats = get_stats(conn)
    total_empresas = conn.execute('SELECT COUNT(*) FROM empresas WHERE ativa=1').fetchone()[0]
    total_urgentes = conn.execute('''
        SELECT COUNT(*) FROM documentos d JOIN empresas e ON e.id=d.empresa_id
        WHERE d.status IN ('VENCIDO','RENOVAR') AND d.vencimento IS NOT NULL AND e.ativa=1
    ''').fetchone()[0]
    def urgentes_de(status, ordem):
        return conn.execute(f'''
            SELECT d.id, d.tipo, d.vencimento, d.status, e.nome as empresa, e.id as emp_id
            FROM documentos d JOIN empresas e ON e.id=d.empresa_id
            WHERE d.status=? AND d.vencimento IS NOT NULL AND e.ativa=1
            ORDER BY d.vencimento {ordem} LIMIT 6
        ''', (status,)).fetchall()
    vencidos_raw = urgentes_de('VENCIDO', 'ASC')   # os mais atrasados primeiro
    renovar_raw  = urgentes_de('RENOVAR', 'ASC')   # os que vencem antes primeiro
    historico = conn.execute('''
        SELECT h.descricao, h.tipo, h.criado_em, e.nome as empresa_nome, u.nome as usuario_nome
        FROM historico h
        LEFT JOIN empresas e ON e.id=h.empresa_id
        LEFT JOIN usuarios u ON u.id=h.usuario_id
        ORDER BY h.criado_em DESC, h.id DESC LIMIT 6
    ''').fetchall()
    hoje = date.today()
    meses = vencimentos_por_mes(conn, hoje)
    conn.close()
    def montar(linhas):
        return [{'id':u['emp_id'],'doc_id':u['id'],'tipo':u['tipo'],'empresa':u['empresa'],
                 'vencimento':u['vencimento'],'status':u['status'],'dias':calcular_dias(u['vencimento'])}
                for u in linhas]
    return render_template('dashboard.html', stats=stats, total_empresas=total_empresas,
                           vencidos=montar(vencidos_raw), renovar=montar(renovar_raw),
                           total_urgentes=total_urgentes,
                           historico=historico, hoje=hoje, meses=meses,
                           max_mes=max([m['total'] for m in meses] + [1]))

# ─── EMPRESAS ─────────────────────────────────────────────────────────────────

@app.route('/empresas')
@login_obrigatorio
def empresas():
    conn = get_connection()
    recalcular_todos(conn)
    cats = conn.execute('SELECT * FROM categorias ORDER BY nome').fetchall()
    dados = []
    for cat in cats:
        emps = conn.execute('''
            SELECT e.id, e.nome, e.cnpj,
                   SUM(CASE WHEN d.status='VENCIDO' THEN 1 ELSE 0 END) as vencidos,
                   SUM(CASE WHEN d.status='RENOVAR' THEN 1 ELSE 0 END) as renovar,
                   SUM(CASE WHEN d.status='OK'      THEN 1 ELSE 0 END) as ok,
                   SUM(CASE WHEN d.status='NÃO TEM' THEN 1 ELSE 0 END) as sem_data
            FROM empresas e
            LEFT JOIN documentos d ON d.empresa_id=e.id
            WHERE e.categoria_id=? AND e.ativa=1
            GROUP BY e.id ORDER BY e.nome
        ''', (cat['id'],)).fetchall()
        dados.append({'categoria':cat,'empresas':emps})
    total = sum(len(d['empresas']) for d in dados)
    conn.close()
    return render_template('empresas.html', dados=dados, total_empresas=total)

@app.route('/empresa/<int:emp_id>')
@login_obrigatorio
def empresa_detalhe(emp_id):
    conn = get_connection()
    recalcular_todos(conn)
    empresa = conn.execute('SELECT e.*, c.nome as categoria FROM empresas e JOIN categorias c ON c.id=e.categoria_id WHERE e.id=?', (emp_id,)).fetchone()
    if not empresa:
        conn.close(); abort(404)
    docs_raw = conn.execute('''SELECT * FROM documentos WHERE empresa_id=?
        ORDER BY CASE status WHEN 'VENCIDO' THEN 0 WHEN 'RENOVAR' THEN 1 WHEN 'NÃO TEM' THEN 2 ELSE 3 END,
                 vencimento IS NULL, vencimento, tipo''', (emp_id,)).fetchall()
    documentos = []
    for doc in docs_raw:
        dias = calcular_dias(doc['vencimento'])
        resps = conn.execute('''SELECT r.id, r.nome, r.email FROM responsaveis r
            JOIN documento_responsavel dr ON dr.responsavel_id=r.id
            WHERE dr.documento_id=? AND r.ativo=1 ORDER BY r.nome''', (doc['id'],)).fetchall()
        documentos.append({'id':doc['id'],'tipo':doc['tipo'],'protocolo':doc['protocolo'],
                           'vencimento':doc['vencimento'],'status':doc['status'],
                           'observacoes':doc['observacoes'] or '','dias':dias,
                           'responsaveis':[dict(r) for r in resps]})
    todos_resp = conn.execute('SELECT * FROM responsaveis WHERE ativo=1 ORDER BY nome').fetchall()
    conn.close()
    resumo = {'VENCIDO': 0, 'RENOVAR': 0, 'OK': 0, 'NÃO TEM': 0}
    for d in documentos:
        resumo[d['status']] = resumo.get(d['status'], 0) + 1
    tipos_faltando = [t for t in TIPOS_DOC if t not in {d['tipo'] for d in documentos}]
    return render_template('empresa_detalhe.html', empresa=empresa, documentos=documentos,
                           todos_responsaveis=todos_resp, resumo=resumo,
                           tipos_doc=TIPOS_DOC, tipos_faltando=tipos_faltando, gerado_em=datetime.now())

# ─── RELATÓRIO ────────────────────────────────────────────────────────────────

@app.route('/relatorio')
@login_obrigatorio
def relatorio():
    filtros = ler_filtros()
    conn = get_connection()
    recalcular_todos(conn)
    docs = consultar_documentos(conn, filtros)
    categorias = conn.execute('SELECT id, nome FROM categorias ORDER BY nome').fetchall()
    lista_empresas = conn.execute('SELECT id, nome FROM empresas WHERE ativa=1 ORDER BY nome').fetchall()
    descricao = descrever_filtros(conn, filtros)
    conn.close()

    grupos, atual = [], None
    for d in docs:
        if not atual or atual['empresa_id'] != d['empresa_id']:
            atual = {'empresa_id': d['empresa_id'], 'empresa': d['empresa'], 'cnpj': d['cnpj'],
                     'categoria': d['categoria'], 'documentos': []}
            grupos.append(atual)
        atual['documentos'].append(d)
    totais = {'VENCIDO': 0, 'RENOVAR': 0, 'OK': 0, 'NÃO TEM': 0}
    for d in docs:
        totais[d['status']] = totais.get(d['status'], 0) + 1

    return render_template('relatorio.html', grupos=grupos, totais=totais, total=len(docs),
                           filtros=filtros, filtros_status=FILTROS_STATUS, descricao=descricao,
                           categorias=categorias, lista_empresas=lista_empresas,
                           gerado_em=datetime.now())

# ─── CADASTROS ────────────────────────────────────────────────────────────────

@app.route('/cadastros')
@login_obrigatorio
def cadastros():
    conn = get_connection()
    categorias = conn.execute('''SELECT c.*, COUNT(e.id) AS total_empresas,
                                        SUM(CASE WHEN e.ativa=1 THEN 1 ELSE 0 END) AS ativas,
                                        SUM(CASE WHEN e.ativa=0 THEN 1 ELSE 0 END) AS inativas
                                 FROM categorias c LEFT JOIN empresas e ON e.categoria_id=c.id
                                 GROUP BY c.id ORDER BY c.nome''').fetchall()
    empresas_ativas = conn.execute('''
        SELECT e.*, c.nome as categoria_nome, COUNT(d.id) as total_docs
        FROM empresas e JOIN categorias c ON c.id=e.categoria_id
        LEFT JOIN documentos d ON d.empresa_id=e.id
        WHERE e.ativa=1 GROUP BY e.id ORDER BY e.nome
    ''').fetchall()
    empresas_inativas = conn.execute('''
        SELECT e.*, c.nome as categoria_nome, COUNT(d.id) as total_docs FROM empresas e
        JOIN categorias c ON c.id=e.categoria_id
        LEFT JOIN documentos d ON d.empresa_id=e.id
        WHERE e.ativa=0 GROUP BY e.id ORDER BY e.nome
    ''').fetchall()
    conn.close()
    aba = request.args.get('aba', 'empresas')
    total_documentos = sum(e['total_docs'] for e in empresas_ativas)
    return render_template('cadastros.html', categorias=categorias,
                           empresas_ativas=empresas_ativas, empresas_inativas=empresas_inativas,
                           tipos_doc=TIPOS_DOC, aba=aba, total_documentos=total_documentos)

@app.route('/empresa/nova', methods=['POST'])
@admin_obrigatorio
def empresa_nova():
    nome   = request.form.get('nome','').strip()
    cnpj   = request.form.get('cnpj','').strip() or None
    cat_id = request.form.get('categoria_id', type=int)
    if not nome or not cat_id:
        flash('Informe o nome e a categoria da empresa.', 'erro')
        return redirect(url_for('cadastros'))
    conn = get_connection()
    if not conn.execute('SELECT 1 FROM categorias WHERE id=?', (cat_id,)).fetchone():
        conn.close(); flash('Categoria não encontrada. Recarregue a página e escolha de novo.', 'erro')
        return redirect(url_for('cadastros'))
    cur = conn.execute('INSERT INTO empresas (nome, cnpj, categoria_id) VALUES (?,?,?)', (nome,cnpj,cat_id))
    registrar_historico(conn, f'Empresa cadastrada: {nome}', empresa_id=cur.lastrowid)
    conn.commit(); conn.close()
    flash(f'Empresa "{nome}" cadastrada.', 'ok')
    return redirect(url_for('cadastros'))

def _nome_empresa(conn, emp_id):
    r = conn.execute('SELECT nome FROM empresas WHERE id=?', (emp_id,)).fetchone()
    if not r: abort(404)
    return r['nome']

@app.route('/empresa/<int:emp_id>/inativar', methods=['POST'])
@admin_obrigatorio
def empresa_inativar(emp_id):
    conn = get_connection()
    nome = _nome_empresa(conn, emp_id)
    conn.execute('UPDATE empresas SET ativa=0 WHERE id=?', (emp_id,))
    registrar_historico(conn, f'Empresa inativada: {nome}', empresa_id=emp_id)
    conn.commit(); conn.close()
    flash(f'"{nome}" foi inativada. Ela continua na aba Inativos.', 'ok')
    return redirect(url_for('cadastros'))

@app.route('/empresa/<int:emp_id>/reativar', methods=['POST'])
@admin_obrigatorio
def empresa_reativar(emp_id):
    conn = get_connection()
    nome = _nome_empresa(conn, emp_id)
    conn.execute('UPDATE empresas SET ativa=1 WHERE id=?', (emp_id,))
    registrar_historico(conn, f'Empresa reativada: {nome}', empresa_id=emp_id)
    conn.commit(); conn.close()
    flash(f'"{nome}" foi reativada.', 'ok')
    return redirect(url_for('cadastros', aba='inativos'))

@app.route('/empresa/<int:emp_id>/excluir', methods=['POST'])
@admin_obrigatorio
def empresa_excluir(emp_id):
    conn = get_connection()
    nome = _nome_empresa(conn, emp_id)
    total = conn.execute('SELECT COUNT(*) FROM documentos WHERE empresa_id=?', (emp_id,)).fetchone()[0]
    if total > 0:
        conn.close(); flash('Não é possível excluir uma empresa com documentos. Inative-a.', 'erro')
        return redirect(request.referrer or url_for('cadastros'))
    conn.execute('DELETE FROM empresas WHERE id=?', (emp_id,))
    registrar_historico(conn, f'Empresa excluída: {nome}')
    conn.commit(); conn.close()
    flash(f'"{nome}" foi excluída.', 'ok')
    return redirect(request.referrer or url_for('cadastros'))

@app.route('/categoria/nova', methods=['POST'])
@admin_obrigatorio
def categoria_nova():
    nome = request.form.get('nome','').strip()
    if nome:
        conn = get_connection()
        try:
            conn.execute('INSERT INTO categorias (nome) VALUES (?)', (nome,))
            conn.commit(); flash(f'Categoria "{nome}" criada.','ok')
        except Exception:
            flash('Já existe uma categoria com esse nome.','erro')
        finally: conn.close()
    return redirect(url_for('cadastros', aba='categorias'))

@app.route('/categoria/<int:cat_id>/excluir', methods=['POST'])
@admin_obrigatorio
def categoria_excluir(cat_id):
    conn = get_connection()
    total = conn.execute('SELECT COUNT(*) FROM empresas WHERE categoria_id=?', (cat_id,)).fetchone()[0]
    if total > 0:
        conn.close(); flash('Essa categoria tem empresas e não pode ser excluída.','erro')
        return redirect(url_for('cadastros', aba='categorias'))
    conn.execute('DELETE FROM categorias WHERE id=?', (cat_id,))
    conn.commit(); conn.close()
    flash('Categoria excluída.','ok')
    return redirect(url_for('cadastros', aba='categorias'))

@app.route('/cadastros/documento/novo', methods=['POST'])
@admin_obrigatorio
def documento_novo():
    emp_id     = request.form.get('empresa_id', type=int)
    tipo       = request.form.get('tipo','').strip()
    protocolo  = request.form.get('protocolo','').strip() or None
    vencimento = data_valida(request.form.get('vencimento'))
    voltar     = request.form.get('voltar')
    destino = url_for('empresa_detalhe', emp_id=emp_id) if voltar == 'empresa' and emp_id else url_for('cadastros', aba='documentos')
    if not emp_id or not tipo:
        flash('Escolha a empresa e o tipo do documento.', 'erro')
        return redirect(destino)
    if request.form.get('vencimento') and not vencimento:
        flash('Data de vencimento inválida.', 'erro')
        return redirect(destino)
    conn = get_connection()
    nome = _nome_empresa(conn, emp_id)
    status = recalcular_status(calcular_dias(vencimento))
    cur = conn.execute('INSERT INTO documentos (empresa_id, tipo, protocolo, vencimento, status) VALUES (?,?,?,?,?)',
                       (emp_id, tipo, protocolo, vencimento, status))
    registrar_historico(conn, f'Documento "{tipo}" cadastrado: {nome}', empresa_id=emp_id, documento_id=cur.lastrowid)
    conn.commit(); conn.close()
    flash(f'"{tipo}" cadastrado em {nome}.','ok')
    return redirect(destino)

# ─── RESPONSÁVEIS ─────────────────────────────────────────────────────────────

@app.route('/responsaveis')
@login_obrigatorio
def responsaveis():
    conn = get_connection()
    lista = conn.execute('''SELECT r.*, (SELECT COUNT(*) FROM documento_responsavel dr
                                         JOIN documentos d ON d.id=dr.documento_id
                                         JOIN empresas e ON e.id=d.empresa_id
                                         WHERE dr.responsavel_id=r.id AND e.ativa=1) AS total_docs
                            FROM responsaveis r WHERE r.ativo=1 ORDER BY r.nome''').fetchall()
    conn.close()
    return render_template('responsaveis.html', responsaveis=lista)

@app.route('/responsaveis/novo', methods=['POST'])
@admin_obrigatorio
def responsavel_novo():
    nome  = request.form.get('nome','').strip()
    email = request.form.get('email','').strip().lower()
    if not nome or not email:
        flash('Informe o nome e o e-mail.', 'erro')
        return redirect(url_for('responsaveis'))
    conn = get_connection()
    existente = conn.execute('SELECT id, ativo FROM responsaveis WHERE lower(email)=?', (email,)).fetchone()
    if existente and existente['ativo']:
        flash('Já existe um responsável com esse e-mail.','erro')
    elif existente:
        conn.execute('UPDATE responsaveis SET nome=?, email=?, ativo=1 WHERE id=?', (nome, email, existente['id']))
        conn.commit(); flash(f'{nome} voltou à lista de responsáveis.','ok')
    else:
        conn.execute('INSERT INTO responsaveis (nome, email) VALUES (?,?)', (nome,email))
        conn.commit(); flash(f'{nome} cadastrado como responsável.','ok')
    conn.close()
    return redirect(url_for('responsaveis'))

@app.route('/responsaveis/<int:resp_id>/excluir', methods=['POST'])
@admin_obrigatorio
def responsavel_excluir(resp_id):
    conn = get_connection()
    r = conn.execute('SELECT nome FROM responsaveis WHERE id=?', (resp_id,)).fetchone()
    if not r: conn.close(); abort(404)
    conn.execute('UPDATE responsaveis SET ativo=0 WHERE id=?', (resp_id,))
    # Os vínculos saem junto: se a pessoa voltar, começa sem documentos atribuídos
    conn.execute('DELETE FROM documento_responsavel WHERE responsavel_id=?', (resp_id,))
    conn.commit(); conn.close()
    flash(f'{r["nome"]} não recebe mais alertas.','ok')
    return redirect(url_for('responsaveis'))

# ─── HISTÓRICO ────────────────────────────────────────────────────────────────

@app.route('/historico')
@login_obrigatorio
def historico():
    per  = 50
    conn = get_connection()
    total = conn.execute('SELECT COUNT(*) FROM historico').fetchone()[0]
    total_pages = max(1, math.ceil(total / per))
    page = request.args.get('page', 1, type=int) or 1
    page = min(max(page, 1), total_pages)
    registros = conn.execute('''
        SELECT h.*, e.nome as empresa_nome, u.nome as usuario_nome
        FROM historico h
        LEFT JOIN empresas e ON e.id=h.empresa_id
        LEFT JOIN usuarios u ON u.id=h.usuario_id
        ORDER BY h.criado_em DESC, h.id DESC LIMIT ? OFFSET ?
    ''', (per, (page-1)*per)).fetchall()
    conn.close()
    return render_template('historico.html', registros=registros, total=total,
                           page=page, total_pages=total_pages)

# ─── USUÁRIOS ─────────────────────────────────────────────────────────────────

@app.route('/usuarios')
@admin_obrigatorio
def usuarios():
    conn = get_connection()
    lista = conn.execute('SELECT * FROM usuarios ORDER BY nome').fetchall()
    conn.close()
    return render_template('usuarios.html', usuarios=lista)

@app.route('/usuarios/novo', methods=['POST'])
@admin_obrigatorio
def usuario_novo():
    nome  = request.form.get('nome','').strip()
    email = request.form.get('email','').strip().lower()
    senha = request.form.get('senha','')
    nivel = request.form.get('nivel','visualizador')
    if nivel not in NIVEIS:
        nivel = 'visualizador'
    if not nome or not email:
        flash('Informe o nome e o e-mail.','erro')
        return redirect(url_for('usuarios'))
    if len(senha) < 6:
        flash('A senha precisa ter pelo menos 6 caracteres.','erro')
        return redirect(url_for('usuarios'))
    conn = get_connection()
    if conn.execute('SELECT 1 FROM usuarios WHERE lower(email)=?', (email,)).fetchone():
        conn.close(); flash('Já existe um usuário com esse e-mail.','erro')
        return redirect(url_for('usuarios'))
    try:
        conn.execute('INSERT INTO usuarios (nome, email, senha, nivel) VALUES (?,?,?,?)',
                     (nome, email, generate_password_hash(senha), nivel))
        conn.commit(); flash(f'{nome} pode entrar no sistema.','ok')
    except Exception:
        flash('Já existe um usuário com esse e-mail.','erro')
    finally: conn.close()
    return redirect(url_for('usuarios'))

@app.route('/usuarios/<int:uid>/excluir', methods=['POST'])
@admin_obrigatorio
def usuario_excluir(uid):
    if uid == session.get('usuario_id'):
        flash('Você não pode excluir o seu próprio usuário.','erro')
        return redirect(url_for('usuarios'))
    conn = get_connection()
    r = conn.execute('SELECT nome FROM usuarios WHERE id=?', (uid,)).fetchone()
    if not r: conn.close(); abort(404)
    conn.execute('DELETE FROM usuarios WHERE id=?', (uid,))
    conn.commit(); conn.close()
    flash(f'{r["nome"]} não tem mais acesso.','ok')
    return redirect(url_for('usuarios'))

@app.route('/usuarios/<int:uid>/redefinir-senha', methods=['POST'])
@admin_obrigatorio
def usuario_redefinir_senha(uid):
    nova = request.form.get('senha_nova','')
    if len(nova) < 6:
        flash('A senha precisa ter pelo menos 6 caracteres.','erro')
        return redirect(url_for('usuarios'))
    conn = get_connection()
    r = conn.execute('SELECT nome FROM usuarios WHERE id=?', (uid,)).fetchone()
    if not r: conn.close(); abort(404)
    conn.execute('UPDATE usuarios SET senha=? WHERE id=?', (generate_password_hash(nova), uid))
    conn.commit(); conn.close()
    flash(f'Senha de {r["nome"]} redefinida.','ok')
    return redirect(url_for('usuarios'))

# ─── CONFIGURAÇÕES ────────────────────────────────────────────────────────────

@app.route('/configuracoes', methods=['GET','POST'])
@admin_obrigatorio
def configuracoes():
    if request.method == 'POST':
        erros = []
        email = request.form.get('email_remetente','').strip()
        if email:
            set_config('email_remetente', email)
        senha_app = request.form.get('email_senha_app','').replace(' ', '')
        if senha_app:
            set_config('email_senha_app', senha_app)
        horario = request.form.get('horario_envio','').strip()
        if horario:
            if re.fullmatch(r'([01]\d|2[0-3]):[0-5]\d', horario): set_config('horario_envio', horario)
            else: erros.append('Horário inválido.')
        url_sistema = request.form.get('url_sistema', '').strip().rstrip('/')
        if url_sistema and not re.match(r'^https?://', url_sistema):
            erros.append('O endereço do sistema deve começar com http:// ou https://.')
        else:
            set_config('url_sistema', url_sistema)
        for campo in ['alerta_dias_90','alerta_dias_30','alerta_dias_7']:
            valor = request.form.get(campo,'').strip()
            if valor:
                if valor.isdigit() and 1 <= int(valor) <= 730: set_config(campo, str(int(valor)))
                else: erros.append('Os prazos de alerta devem ser números entre 1 e 730.')
        horario = get_config('horario_envio') or '08:00'
        hora, minuto = horario.split(':')
        if scheduler:
            try:
                scheduler.reschedule_job('verificacao_diaria', trigger='cron', hour=int(hora), minute=int(minuto))
            except Exception:
                app.logger.exception('Não foi possível reagendar a verificação diária')
        if erros:
            for e in dict.fromkeys(erros): flash(e, 'erro')
        else:
            flash('Configurações salvas.','ok')
        return redirect(url_for('configuracoes'))
    cfg = {
        'email_remetente': get_config('email_remetente') or '',
        'url_sistema':     get_config('url_sistema') or '',
        'tem_senha_app':   bool(get_config('email_senha_app')),
        'horario_envio':   get_config('horario_envio') or '08:00',
        'alerta_dias_90':  get_config('alerta_dias_90') or '90',
        'alerta_dias_30':  get_config('alerta_dias_30') or '30',
        'alerta_dias_7':   get_config('alerta_dias_7') or '7',
    }
    return render_template('configuracoes.html', cfg=cfg)

@app.route('/api/testar-email', methods=['POST'])
@admin_obrigatorio
def testar_email():
    from notificacoes import enviar_email
    if MODO_DEMO:
        return jsonify({'ok': False, 'demo': True, 'erro': 'No modo demonstração nenhum e-mail é enviado.'})
    if not get_config('email_remetente') or not get_config('email_senha_app'):
        return jsonify({'ok': False, 'erro': 'Preencha o e-mail remetente e a senha de app, salve e tente de novo.'})
    conn = get_connection()
    u = conn.execute('SELECT email, nome FROM usuarios WHERE id=?', (session['usuario_id'],)).fetchone()
    conn.close()
    ok = enviar_email(u['email'], u['nome'], 'AlertSignal: teste de e-mail',
                      '<p>Se você recebeu este e-mail, a configuração está correta.</p>')
    if ok:
        return jsonify({'ok': True, 'mensagem': f'E-mail de teste enviado para {u["email"]}.'})
    return jsonify({'ok': False, 'erro': 'O Gmail recusou o envio. Confira o e-mail remetente e a senha de app.'})

# ─── API DOCUMENTOS ───────────────────────────────────────────────────────────

@app.route('/api/doc/<int:doc_id>/responsavel', methods=['POST'])
@admin_obrigatorio
def vincular_responsavel(doc_id):
    data = request.get_json(silent=True) or {}
    resp_id = data.get('responsavel_id')
    conn = get_connection()
    doc = dados_documento(conn, doc_id)
    resp = conn.execute('SELECT id, nome FROM responsaveis WHERE id=? AND ativo=1', (resp_id,)).fetchone()
    if not doc or not resp:
        conn.close(); return jsonify({'ok': False, 'erro': 'Documento ou responsável não encontrado.'}), 404
    conn.execute('INSERT OR IGNORE INTO documento_responsavel (documento_id, responsavel_id) VALUES (?,?)', (doc_id, resp['id']))
    registrar_historico(conn, f'{resp["nome"]} agora responde por {doc["tipo"]}, {doc["empresa"]}',
                        empresa_id=doc['empresa_id'], documento_id=doc_id)
    conn.commit(); conn.close()
    return jsonify({'ok': True, 'id': resp['id'], 'nome': resp['nome']})

@app.route('/api/doc/<int:doc_id>/responsavel/<int:resp_id>', methods=['DELETE'])
@admin_obrigatorio
def desvincular_responsavel(doc_id, resp_id):
    conn = get_connection()
    doc = dados_documento(conn, doc_id)
    resp = conn.execute('SELECT nome FROM responsaveis WHERE id=?', (resp_id,)).fetchone()
    conn.execute('DELETE FROM documento_responsavel WHERE documento_id=? AND responsavel_id=?', (doc_id, resp_id))
    if doc and resp:
        registrar_historico(conn, f'{resp["nome"]} deixou de responder por {doc["tipo"]}, {doc["empresa"]}',
                            empresa_id=doc['empresa_id'], documento_id=doc_id)
    conn.commit(); conn.close()
    return jsonify({'ok': True})

@app.route('/api/doc/<int:doc_id>/protocolo', methods=['POST'])
@admin_obrigatorio
def atualizar_protocolo(doc_id):
    data = request.get_json(silent=True) or {}
    protocolo = (data.get('protocolo') or '').strip() or None
    conn = get_connection()
    doc = dados_documento(conn, doc_id)
    if not doc:
        conn.close(); return jsonify({'ok': False, 'erro': 'Documento não encontrado.'}), 404
    conn.execute('UPDATE documentos SET protocolo=? WHERE id=?', (protocolo, doc_id))
    registrar_historico(conn, f'Protocolo atualizado: {doc["tipo"]}, {doc["empresa"]}',
                        empresa_id=doc['empresa_id'], documento_id=doc_id)
    conn.commit(); conn.close()
    return jsonify({'ok': True, 'protocolo': protocolo})

@app.route('/api/doc/<int:doc_id>/editar', methods=['POST'])
@admin_obrigatorio
def editar_documento(doc_id):
    data = request.get_json(silent=True) or {}
    protocolo   = (data.get('protocolo') or '').strip() or None
    vencimento  = data_valida(data.get('vencimento'))
    observacoes = (data.get('observacoes') or '').strip() or None
    if data.get('vencimento') and not vencimento:
        return jsonify({'ok': False, 'erro': 'Data de vencimento inválida.'}), 400
    conn = get_connection()
    doc = dados_documento(conn, doc_id)
    if not doc:
        conn.close(); return jsonify({'ok': False, 'erro': 'Documento não encontrado.'}), 404
    status = recalcular_status(calcular_dias(vencimento))
    conn.execute('UPDATE documentos SET protocolo=?, vencimento=?, observacoes=?, status=? WHERE id=?',
                 (protocolo, vencimento, observacoes, status, doc_id))
    registrar_historico(conn, f'Documento "{doc["tipo"]}" editado: {doc["empresa"]}',
                        empresa_id=doc['empresa_id'], documento_id=doc_id)
    conn.commit(); conn.close()
    return jsonify({'ok': True})

@app.route('/api/doc/<int:doc_id>/renovar', methods=['POST'])
@admin_obrigatorio
def renovar_documento(doc_id):
    data = request.get_json(silent=True) or {}
    vencimento = data_valida(data.get('vencimento'))
    if not vencimento:
        return jsonify({'ok': False, 'erro': 'Informe a nova data de vencimento.'}), 400
    if vencimento <= date.today().isoformat():
        return jsonify({'ok': False, 'erro': 'A nova data precisa ser depois de hoje.'}), 400
    conn = get_connection()
    doc = dados_documento(conn, doc_id)
    if not doc:
        conn.close(); return jsonify({'ok': False, 'erro': 'Documento não encontrado.'}), 404
    status = recalcular_status(calcular_dias(vencimento))
    conn.execute('UPDATE documentos SET vencimento=?, status=? WHERE id=?', (vencimento, status, doc_id))
    registrar_historico(conn, f'Documento "{doc["tipo"]}" renovado: {doc["empresa"]}',
                        empresa_id=doc['empresa_id'], documento_id=doc_id)
    conn.commit(); conn.close()
    return jsonify({'ok': True})

@app.route('/api/doc/<int:doc_id>/excluir', methods=['DELETE'])
@admin_obrigatorio
def excluir_documento(doc_id):
    conn = get_connection()
    doc = dados_documento(conn, doc_id)
    if not doc:
        conn.close(); return jsonify({'ok': False, 'erro': 'Documento não encontrado.'}), 404
    conn.execute('DELETE FROM documento_responsavel WHERE documento_id=?', (doc_id,))
    conn.execute('DELETE FROM documentos WHERE id=?', (doc_id,))
    registrar_historico(conn, f'Documento "{doc["tipo"]}" excluído: {doc["empresa"]}', empresa_id=doc['empresa_id'])
    conn.commit(); conn.close()
    return jsonify({'ok': True})

@app.route('/api/doc/<int:doc_id>/reenviar-alerta', methods=['POST'])
@admin_obrigatorio
def reenviar_alerta(doc_id):
    from notificacoes import montar_html, montar_texto, enviar_email, nivel_alerta, registrar_historico as reg_hist
    conn = get_connection()
    doc = conn.execute('''
        SELECT d.*, e.nome as empresa, e.ativa FROM documentos d
        JOIN empresas e ON e.id=d.empresa_id WHERE d.id=?
    ''', (doc_id,)).fetchone()
    resps = conn.execute('''SELECT r.nome, r.email FROM responsaveis r
        JOIN documento_responsavel dr ON dr.responsavel_id=r.id
        WHERE dr.documento_id=? AND r.ativo=1''', (doc_id,)).fetchall()
    conn.close()
    if not doc:
        return jsonify({'ok': False, 'erro': 'Documento não encontrado.'}), 404
    if not doc['ativa']:
        return jsonify({'ok': False, 'erro': 'A empresa está inativa e não recebe alertas.'}), 400
    if not doc['vencimento']:
        return jsonify({'ok': False, 'erro': 'Este documento não tem vencimento.'}), 400
    if not resps:
        return jsonify({'ok': False, 'erro': 'Este documento ainda não tem responsável.'})
    if MODO_DEMO:
        return jsonify({'ok': False, 'demo': True, 'erro': 'No modo demonstração nenhum alerta é enviado.'})
    dias = calcular_dias(doc['vencimento'])
    enviados = 0
    for r in resps:
        itens = [{'empresa':doc['empresa'],'documento':doc['tipo'],'vencimento':doc['vencimento'],'dias':dias,
                  'nivel':nivel_alerta(dias) or 'antecipado'}]
        corpo = montar_html(r['nome'], itens, get_config('url_sistema'))
        texto = montar_texto(r['nome'], itens, get_config('url_sistema'))
        if enviar_email(r['email'], r['nome'], f'AlertSignal: {doc["tipo"]} de {doc["empresa"]}', corpo, texto):
            enviados += 1
    if not enviados:
        return jsonify({'ok': False, 'erro': 'Nenhum e-mail foi enviado. Confira as configurações de e-mail.'})
    reg_hist(f'Alerta reenviado manualmente: {doc["tipo"]}, {doc["empresa"]}',
             empresa_id=doc['empresa_id'], documento_id=doc_id)
    return jsonify({'ok': True, 'mensagem': f'Alerta enviado para {enviados} responsável(is).'})

# ─── EXPORTAR ─────────────────────────────────────────────────────────────────

COLUNAS_EXPORTACAO = ['Empresa', 'CNPJ', 'Categoria', 'Documento', 'Protocolo', 'Vencimento',
                      'Dias para vencer', 'Status', 'Responsáveis', 'Observações']

@app.route('/exportar')
@login_obrigatorio
def exportar():
    import csv
    import io
    from flask import send_file
    filtros = ler_filtros()
    conn = get_connection()
    recalcular_todos(conn)
    docs = consultar_documentos(conn, filtros)
    descricao = descrever_filtros(conn, filtros)
    nome_empresa = _nome_empresa(conn, filtros['empresa']) if filtros.get('empresa') else None
    conn.close()

    base = f"alertsignal_{_slug(nome_empresa) if nome_empresa else 'documentos'}_{date.today().isoformat()}"
    fmt = request.args.get('fmt', 'xlsx')

    if fmt == 'csv':
        # Ponto e vírgula e datas dd/mm/aaaa: é o que o Excel em português abre
        # direto em colunas. O BOM avisa que o arquivo é UTF-8.
        texto = io.StringIO()
        escritor = csv.writer(texto, delimiter=';', lineterminator='\r\n')
        escritor.writerow(COLUNAS_EXPORTACAO)
        for d in docs:
            escritor.writerow([d['empresa'], d['cnpj'] or '', d['categoria'], d['tipo'], d['protocolo'] or '',
                               formatar_data(d['vencimento']) if d['vencimento'] else '',
                               '' if d['dias'] is None else d['dias'], status_rotulo(d['status']),
                               d['responsaveis'] or '', d['observacoes'] or ''])
        buf = io.BytesIO(texto.getvalue().encode('utf-8-sig'))
        return send_file(buf, mimetype='text/csv', as_attachment=True, download_name=f'{base}.csv')

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Documentos'
    wb.properties.title = 'AlertSignal: documentos'
    wb.properties.creator = 'AlertSignal'

    fonte = 'Calibri'
    ws['A1'] = 'AlertSignal · Documentos e vencimentos'
    ws['A1'].font = Font(name=fonte, bold=True, size=14, color='111113')
    ws['A2'] = f'{descricao} · gerado em {datetime.now().strftime("%d/%m/%Y às %H:%M")} · {len(docs)} documento(s)'
    ws['A2'].font = Font(name=fonte, size=10, color='5B5B63')
    totais = {'VENCIDO': 0, 'RENOVAR': 0, 'OK': 0, 'NÃO TEM': 0}
    for d in docs:
        totais[d['status']] = totais.get(d['status'], 0) + 1
    ws['A3'] = (f"Vencidos: {totais['VENCIDO']} · A renovar: {totais['RENOVAR']} · Em dia: {totais['OK']}"
                f" · Sem data: {totais['NÃO TEM']} · Dias negativos indicam documento já vencido")
    ws['A3'].font = Font(name=fonte, size=10, color='5B5B63')
    ws.row_dimensions[1].height = 22

    LINHA_CAB = 5
    borda = Border(bottom=Side(style='thin', color='E2E2E6'))
    for col, titulo in enumerate(COLUNAS_EXPORTACAO, 1):
        cell = ws.cell(row=LINHA_CAB, column=col, value=titulo)
        cell.font = Font(name=fonte, bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor='1F1F23')
        cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[LINHA_CAB].height = 24

    # Texto e fundo por status, com contraste suficiente para leitura e impressão
    CORES_STATUS = {
        'VENCIDO': ('B42318', 'FDECEC'),
        'RENOVAR': ('93580B', 'FEF4E2'),
        'OK':      ('067647', 'E8F7EF'),
        'NÃO TEM': ('475467', 'F2F4F7'),
    }
    zebra = PatternFill('solid', fgColor='F7F7F9')
    for i, d in enumerate(docs):
        linha = LINHA_CAB + 1 + i
        venc = datetime.strptime(d['vencimento'], '%Y-%m-%d') if d['vencimento'] else None
        valores = [d['empresa'], d['cnpj'] or '', d['categoria'], d['tipo'], d['protocolo'] or '',
                   venc, d['dias'], status_rotulo(d['status']), d['responsaveis'] or '', d['observacoes'] or '']
        for col, valor in enumerate(valores, 1):
            cell = ws.cell(row=linha, column=col, value=valor)
            cell.font = Font(name=fonte, size=10, color='1F1F23')
            cell.alignment = Alignment(vertical='top', wrap_text=True, indent=1)
            cell.border = borda
            if i % 2:
                cell.fill = zebra
        ws.cell(row=linha, column=6).number_format = 'DD/MM/YYYY'
        dias = ws.cell(row=linha, column=7)
        dias.number_format = '0;[Red]-0'
        dias.alignment = Alignment(horizontal='right', vertical='top', indent=1)
        texto_cor, fundo = CORES_STATUS.get(d['status'], CORES_STATUS['NÃO TEM'])
        st = ws.cell(row=linha, column=8)
        st.font = Font(name=fonte, size=10, bold=True, color=texto_cor)
        st.fill = PatternFill('solid', fgColor=fundo)

    larguras = [32, 20, 14, 26, 18, 13, 17, 11, 26, 34]
    for i, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ultima = LINHA_CAB + max(len(docs), 1)
    ws.auto_filter.ref = f'A{LINHA_CAB}:{get_column_letter(len(COLUNAS_EXPORTACAO))}{ultima}'
    ws.freeze_panes = f'A{LINHA_CAB + 1}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{LINHA_CAB}:{LINHA_CAB}'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'{base}.xlsx')

# ─── AGENDADOR ────────────────────────────────────────────────────────────────

def iniciar_agendador():
    global scheduler
    scheduler = BackgroundScheduler(timezone='America/Bahia')
    horario = get_config('horario_envio') or '08:00'
    hora, minuto = horario.split(':')
    scheduler.add_job(executar_verificacao_diaria, trigger='cron',
                      hour=int(hora), minute=int(minuto),
                      id='verificacao_diaria', replace_existing=True)
    scheduler.start()

# ─── INICIALIZAÇÃO ────────────────────────────────────────────────────────────

def _argumento(nome, padrao=None):
    if nome in sys.argv:
        i = sys.argv.index(nome)
        if i + 1 < len(sys.argv):
            return sys.argv[i + 1]
    return padrao

if __name__ == '__main__':
    porta = int(_argumento('--porta', os.environ.get('PORT', 5000)))
    if MODO_DEMO:
        from demo_seed import popular, SENHA_DEMO
        database.usar_banco(os.environ.get('ALERTSIGNAL_BANCO_DEMO')
                            or os.path.join(database.PASTA, 'demo.db'))
        popular()
        print(f'\n AlertSignal (DEMONSTRAÇÃO) rodando em http://localhost:{porta}')
        print(f' Dados fictícios em {database.DB_PATH}. Nenhum e-mail é enviado.')
        print(f' Login: admin@alertsignal.com | Senha: {SENHA_DEMO}\n')
        app.run(host='127.0.0.1', port=porta, debug=False)
    else:
        init_db()
        inserir_configuracoes_padrao()
        criar_admin_padrao()
        xlsx = os.path.join(os.path.dirname(__file__), 'ALVARAS_GRUPO_ZEN.xlsx')
        if os.path.exists(xlsx):
            importar(xlsx)
        iniciar_agendador()
        atexit.register(lambda: scheduler.shutdown(wait=False))
        print(f'\n AlertSignal rodando em http://localhost:{porta}')
        print(' Login: admin@grupozen.com.br | Senha: zen2024\n')
        app.run(host='0.0.0.0', port=porta, debug=False)
